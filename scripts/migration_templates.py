"""Menghasilkan template Excel migrasi data Fase 1 (PRD F0-MIG-01 s.d. F0-MIG-04).

Template yang sama nantinya dihasilkan modul `dataexchange` (Technical Guideline 5.15), sehingga kolom di
sini adalah kontrak import. Jalankan:  python scripts/migration_templates.py docs/migrasi
"""

from __future__ import annotations

import sys
from dataclasses import dataclass, field
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

HEADER_FILL = PatternFill("solid", fgColor="C2410C")
OPTIONAL_FILL = PatternFill("solid", fgColor="6B7280")
EXAMPLE_FONT = Font(italic=True, color="6B7280")
TITLE_FONT = Font(bold=True, size=14)


@dataclass
class Col:
    name: str
    required: bool
    fmt: str
    rule: str
    example: str
    choices: list[str] = field(default_factory=list)
    width: int = 18


@dataclass
class Sheet:
    title: str
    purpose: str
    pic: str
    reference: str
    columns: list[Col]
    examples: list[list[str]]


def c(name, required, fmt, rule, example, choices=None, width=18):
    return Col(name, required, fmt, rule, example, choices or [], width)


TEXT = "Teks"
DATE = "Tanggal YYYY-MM-DD"
MONEY = "Angka rupiah tanpa titik/Rp, desimal pakai titik (mis. 1250000 atau 1250.50)"
QTY = "Angka, maks. 3 desimal (pakai titik)"
PCT = "Angka persen (mis. 20 untuk 20%)"
YESNO = ["Ya", "Tidak"]

MASTER = [
    Sheet(
        "Produk",
        "Produk jadi (FG) dan bundel virtual yang dijual",
        "Logistik & Sales",
        "MD-01, MD-03",
        [
            c("SKU", True, "Huruf besar/angka/tanda minus, maks. 40", "Unik, tidak boleh dipakai ulang", "BLS-CHO-TB30", width=16),
            c("Nama Produk", True, TEXT, "Maks. 160 karakter", "Balans Meal Shake Coklat", width=32),
            c("Varian/Rasa", False, TEXT, "", "Coklat"),
            c("Tipe Kemasan", True, "Pilihan", "Tube / Pouch / Roll Sachet", "Tube", ["Tube", "Pouch", "Roll Sachet"]),
            c("Satuan", True, TEXT, "pcs, box, dll.", "pcs", width=10),
            c("Barcode", False, "Angka", "Pisahkan dengan koma jika lebih dari satu", "8991234567890"),
            c("Umur Simpan (hari)", True, "Angka bulat", "> 0", "540"),
            c("Kategori", False, TEXT, "", "Meal Shake"),
            c("Bundel Virtual", True, "Pilihan", "Ya jika SKU ini paket berisi SKU lain (isi di sheet Bundling)", "Tidak", YESNO),
            c("Stok Minimum", False, QTY, "Untuk peringatan stok minimum", "100"),
        ],
        [["BLS-CHO-TB30", "Balans Meal Shake Coklat 30 g", "Coklat", "Tube", "pcs", "8991234567890", "540", "Meal Shake", "Tidak", "100"],
         ["BLS-BUNDLE-3", "Paket Coba 3 Rasa", "", "Tube", "box", "", "540", "Paket", "Ya", ""]],
    ),
    Sheet(
        "Bundling",
        "Komponen setiap bundel virtual",
        "Sales",
        "MD-03",
        [
            c("SKU Bundel", True, TEXT, "Harus ada di sheet Produk dengan Bundel Virtual = Ya", "BLS-BUNDLE-3"),
            c("SKU Komponen", True, TEXT, "Produk FG non-bundel", "BLS-CHO-TB30"),
            c("Qty per Bundel", True, QTY, "> 0", "1"),
        ],
        [["BLS-BUNDLE-3", "BLS-CHO-TB30", "1"], ["BLS-BUNDLE-3", "BLS-VAN-TB30", "1"], ["BLS-BUNDLE-3", "BLS-STR-TB30", "1"]],
    ),
    Sheet(
        "Packaging & Material",
        "Packaging (Tube, Pouch, Roll Sachet) dan bahan reprocess",
        "Logistik",
        "MD-02, U-01",
        [
            c("Kode", True, TEXT, "Unik, tidak boleh dipakai ulang", "PKG-TUBE-30"),
            c("Nama", True, TEXT, "", "Tube kosong 30 g", width=28),
            c("Jenis", True, "Pilihan", "Packaging / Material (bahan reprocess)", "Packaging", ["Packaging", "Material"]),
            c("Tipe Kemasan", False, "Pilihan", "Untuk packaging", "Tube", ["Tube", "Pouch", "Roll Sachet"]),
            c("Satuan", True, TEXT, "", "pcs", width=10),
            c("Pakai Batch", True, "Pilihan", "Ya jika packaging dilacak per batch", "Tidak", YESNO),
            c("Stok Minimum", False, QTY, "", "5000"),
        ],
        [["PKG-TUBE-30", "Tube kosong 30 g", "Packaging", "Tube", "pcs", "Tidak", "5000"]],
    ),
    Sheet(
        "Pelanggan",
        "Pelanggan WA/IG, B2B/corporate, dan pelanggan lain",
        "Sales (CS)",
        "MD-06, Catatan #8",
        [
            c("Kode Pelanggan", False, TEXT, "Kosongkan untuk dibuat otomatis", ""),
            c("Nama Lengkap", True, TEXT, "", "Siti Rahma", width=24),
            c("Nomor WA/HP", True, "08xxxxxxxxxx atau 628xxxxxxxxxx", "Unik; dipakai mencegah duplikat", "081234567890"),
            c("Email", False, "Email", "", "siti@example.com", width=24),
            c("Tipe Pelanggan", True, "Pilihan", "", "Retail", ["Retail", "B2B", "Karyawan"]),
            c("Alamat Pengiriman", True, TEXT, "Jalan, nomor, RT/RW", "Jl. Melati No. 5", width=32),
            c("Kecamatan", True, TEXT, "", "Tebet"),
            c("Kota/Kabupaten", True, TEXT, "", "Jakarta Selatan"),
            c("Provinsi", True, TEXT, "", "DKI Jakarta"),
            c("Kode Pos", False, "5 digit", "", "12810", width=10),
            c("Nomor Karyawan", False, TEXT, "Isi jika pelanggan adalah karyawan (harga karyawan)", ""),
        ],
        [["", "Siti Rahma", "081234567890", "siti@example.com", "Retail", "Jl. Melati No. 5", "Tebet", "Jakarta Selatan", "DKI Jakarta", "12810", ""]],
    ),
    Sheet(
        "Mitra Konsinyasi",
        "Partner Modern Trade (beli putus) dan konsinyasi",
        "Sales MT",
        "MD-07",
        [
            c("Kode Mitra", True, TEXT, "Unik, tidak boleh dipakai ulang", "MT-SINARJAYA"),
            c("Nama Mitra", True, TEXT, "", "Toko Sinar Jaya", width=24),
            c("Jenis", True, "Pilihan", "", "Konsinyasi", ["Modern Trade", "Konsinyasi", "Keduanya"]),
            c("NPWP", False, "15/16 digit", "", "", width=20),
            c("Alamat", True, TEXT, "", "Jl. Raya Bogor 10, Depok", width=32),
            c("Nama PIC", True, TEXT, "", "Budi"),
            c("HP PIC", True, "08xxxxxxxxxx", "", "081298765432"),
            c("Sales Penanggung Jawab (email)", True, "Email user ERP", "", "sales.mt@findyourbalans.com", width=28),
        ],
        [["MT-SINARJAYA", "Toko Sinar Jaya", "Konsinyasi", "", "Jl. Raya Bogor 10, Depok", "Budi", "081298765432", "sales.mt@findyourbalans.com"]],
    ),
    Sheet(
        "PKS",
        "Perjanjian Kerja Sama (PKS) mitra; wajib untuk konsinyasi",
        "Sales MT",
        "MD-07, BRD 5.5",
        [
            c("Kode Mitra", True, TEXT, "Ada di sheet Mitra Konsinyasi", "MT-SINARJAYA"),
            c("Nomor PKS", True, TEXT, "", "PKS/BLS/2026/014"),
            c("Jenis PKS", True, "Pilihan", "", "Konsinyasi", ["Konsinyasi", "Modern Trade"]),
            c("Berlaku Dari", True, DATE, "", "2026-01-01", width=14),
            c("Berlaku Sampai", True, DATE, ">= Berlaku Dari", "2027-12-31", width=14),
            c("Termin Pembayaran (hari)", True, "Angka bulat", "", "30"),
            c("Barang Hilang Ditanggung", True, "Pilihan", "Default: Mitra", "Mitra", ["Mitra", "Balans"]),
            c("Rusak oleh Mitra Ditanggung", True, "Pilihan", "Default: Mitra", "Mitra", ["Mitra", "Balans"]),
            c("Cacat Produk Ditanggung", True, "Pilihan", "Default: Balans", "Balans", ["Mitra", "Balans"]),
            c("Mendekati Kedaluwarsa Ditanggung", True, "Pilihan", "Default: Balans", "Balans", ["Mitra", "Balans"]),
            c("Nama File Dokumen PKS", False, "Nama file PDF", "Lampirkan file PDF bersama template", "PKS-SINARJAYA.pdf", width=24),
        ],
        [["MT-SINARJAYA", "PKS/BLS/2026/014", "Konsinyasi", "2026-01-01", "2027-12-31", "30", "Mitra", "Mitra", "Balans", "Balans", "PKS-SINARJAYA.pdf"]],
    ),
    Sheet(
        "Pemasok & Maklon",
        "Pemasok packaging/barang dan maklon",
        "Logistik & F&A",
        "MD-08",
        [
            c("Kode Pemasok", True, TEXT, "Unik, tidak boleh dipakai ulang", "SUP-KEMASJAYA"),
            c("Nama", True, TEXT, "", "PT Kemas Jaya", width=24),
            c("Jenis", True, "Pilihan", "", "Pemasok", ["Pemasok", "Maklon"]),
            c("Badan Usaha", True, "Pilihan", "Untuk penentuan PPh 23", "Ya", YESNO),
            c("NPWP", False, "15/16 digit", "", "", width=20),
            c("Alamat", True, TEXT, "", "Kawasan Industri Cikarang", width=32),
            c("Termin Pembayaran (hari)", False, "Angka bulat", "", "30"),
            c("Bank", False, TEXT, "", "BCA"),
            c("Nomor Rekening", False, "Angka", "", ""),
            c("Nama Pemilik Rekening", False, TEXT, "", "PT Kemas Jaya", width=24),
        ],
        [["SUP-KEMASJAYA", "PT Kemas Jaya", "Pemasok", "Ya", "", "Kawasan Industri Cikarang", "30", "BCA", "", "PT Kemas Jaya"],
         ["MKL-NUTRIPRO", "PT Nutri Pro Maklon", "Maklon", "Ya", "", "Sentul, Bogor", "14", "Mandiri", "", "PT Nutri Pro Maklon"]],
    ),
    Sheet(
        "Gudang & Lokasi",
        "Gudang fisik dan lokasi stok, termasuk lokasi eksternal",
        "Logistik",
        "MD-09, U-08",
        [
            c("Kode Gudang", False, TEXT, "Kosong untuk lokasi eksternal", "GU"),
            c("Nama Gudang", False, TEXT, "", "Gudang Utama"),
            c("Alamat Gudang", False, TEXT, "", "", width=28),
            c("Kode Lokasi", True, TEXT, "Unik", "GU-INBOUND"),
            c("Nama Lokasi", True, TEXT, "", "Gudang Inbound", width=24),
            c("Tipe Lokasi", True, "Pilihan", "", "Inbound",
              ["Inbound", "FG Non-Ecommerce", "E-commerce Stock", "Gudang Material", "Gudang Retur", "Maklon", "Konsinyasi", "Buffer Vending", "Mesin Vending"]),
            c("Kode Pihak Eksternal", False, TEXT, "Kode maklon/mitra untuk lokasi eksternal", ""),
        ],
        [["GU", "Gudang Utama", "", "GU-INBOUND", "Gudang Inbound", "Inbound", ""],
         ["GU", "Gudang Utama", "", "GU-FG", "FG Non-Ecommerce", "FG Non-Ecommerce", ""],
         ["", "", "", "KSN-SINARJAYA", "Konsinyasi Toko Sinar Jaya", "Konsinyasi", "MT-SINARJAYA"],
         ["", "", "", "MKL-NUTRIPRO", "Lokasi Maklon Nutri Pro", "Maklon", "MKL-NUTRIPRO"]],
    ),
    Sheet(
        "Karyawan",
        "Data dasar karyawan untuk harga karyawan dan akun pengguna",
        "HRGA",
        "HR-01, HR-03",
        [
            c("Nomor Karyawan", True, TEXT, "Unik, tidak boleh dipakai ulang", "EMP-0012"),
            c("Nama Lengkap", True, TEXT, "", "Andi Wijaya", width=24),
            c("Divisi", True, "Pilihan", "", "Sales", ["Logistik", "Finance & Accounting", "Sales", "Marketing", "HRGA", "Manajemen"]),
            c("Jabatan", True, TEXT, "", "Sales MT"),
            c("Status", True, "Pilihan", "", "Aktif", ["Aktif", "Resign"]),
            c("Tanggal Masuk", True, DATE, "", "2024-03-01", width=14),
            c("Tanggal Resign", False, DATE, "Wajib jika Status = Resign", "", width=14),
            c("Email Kantor", False, "Email", "Untuk akun ERP", "andi@findyourbalans.com", width=26),
            c("Nomor HP", False, "08xxxxxxxxxx", "Untuk harga karyawan via CS", "081311112222"),
        ],
        [["EMP-0012", "Andi Wijaya", "Sales", "Sales MT", "Aktif", "2024-03-01", "", "andi@findyourbalans.com", "081311112222"]],
    ),
]

PRICE = [
    Sheet(
        "Harga Dasar & HET",
        "Harga dasar dan Harga Eceran Tertinggi per SKU",
        "Sales & F&A",
        "MD-04, BRD 5.18",
        [
            c("SKU", True, TEXT, "Ada di master produk", "BLS-CHO-TB30"),
            c("Harga Dasar", True, MONEY, "> 0", "25000"),
            c("HET", True, MONEY, ">= Harga Dasar", "35000"),
            c("Berlaku Dari", True, DATE, "", "2027-01-01", width=14),
        ],
        [["BLS-CHO-TB30", "25000", "35000", "2027-01-01"]],
    ),
    Sheet(
        "Daftar Harga",
        "Price list per channel, per mitra, per level reseller, karyawan, vending",
        "Sales",
        "MD-04",
        [
            c("Kode Daftar Harga", True, TEXT, "Unik", "PL-WEBSITE"),
            c("Nama", True, TEXT, "", "Harga Website", width=26),
            c("Jenis", True, "Pilihan", "", "Website",
              ["Modern Trade", "Konsinyasi", "Marketplace", "Website", "WA/IG", "Reseller", "Karyawan", "Vending"]),
            c("Kode Mitra", False, TEXT, "Untuk harga per mitra MT/konsinyasi", ""),
            c("Level Reseller", False, "Pilihan", "Untuk jenis Reseller", "", ["Bronze", "Silver", "Gold"]),
            c("Dasar Harga", True, "Pilihan", "Tetap per SKU, atau persen dari harga website", "Tetap", ["Tetap", "Persen dari Website"]),
            c("Persen", False, PCT, "Untuk dasar persen; diskon karyawan awal 20", ""),
        ],
        [["PL-WEBSITE", "Harga Website", "Website", "", "", "Tetap", ""],
         ["PL-KARYAWAN", "Harga Karyawan", "Karyawan", "", "", "Persen dari Website", "20"]],
    ),
    Sheet(
        "Item Daftar Harga",
        "Harga per SKU untuk daftar harga berdasar Tetap",
        "Sales",
        "MD-04",
        [
            c("Kode Daftar Harga", True, TEXT, "Ada di sheet Daftar Harga", "PL-WEBSITE"),
            c("SKU", True, TEXT, "", "BLS-CHO-TB30"),
            c("Harga", True, MONEY, "> 0", "32000"),
            c("Qty Minimum", False, QTY, "Isi untuk harga bertingkat; kosong = 1", ""),
            c("Berlaku Dari", True, DATE, "", "2027-01-01", width=14),
            c("Berlaku Sampai", False, DATE, "Kosong = tanpa batas", "", width=14),
        ],
        [["PL-WEBSITE", "BLS-CHO-TB30", "32000", "", "2027-01-01", ""],
         ["PL-WEBSITE", "BLS-CHO-TB30", "30000", "12", "2027-01-01", ""]],
    ),
    Sheet(
        "Harga Khusus",
        "Harga khusus per pelanggan atau mitra (prioritas tertinggi)",
        "Sales",
        "MD-04, BRD 5.18",
        [
            c("Kode Pelanggan/Mitra", True, TEXT, "Kode di master pelanggan atau mitra", "MT-SINARJAYA"),
            c("SKU", True, TEXT, "", "BLS-CHO-TB30"),
            c("Harga", True, MONEY, "> 0", "27500"),
            c("Berlaku Dari", True, DATE, "", "2027-01-01", width=14),
            c("Berlaku Sampai", False, DATE, "", "2027-12-31", width=14),
        ],
        [["MT-SINARJAYA", "BLS-CHO-TB30", "27500", "2027-01-01", "2027-12-31"]],
    ),
]

OPENING = [
    Sheet(
        "Saldo Stok",
        "Saldo stok per SKU per batch per lokasi pada tanggal cut-over (hasil stock count fisik)",
        "Logistik & F&A",
        "U-10, A-36",
        [
            c("Kode Lokasi", True, TEXT, "Ada di Gudang & Lokasi, termasuk maklon & konsinyasi", "GU-FG"),
            c("SKU/Kode", True, TEXT, "Produk atau packaging", "BLS-CHO-TB30"),
            c("Nomor Batch", True, TEXT, "Wajib untuk produk ber-batch", "B2312"),
            c("Tanggal Produksi", False, DATE, "", "2026-06-01", width=14),
            c("Tanggal Kedaluwarsa", True, DATE, "Wajib untuk FG", "2027-12-12", width=14),
            c("Qty", True, QTY, "> 0", "1200"),
            c("HPP per Unit", True, "Angka, maks. 4 desimal", "> 0; dari pembukuan lama", "14250.5"),
        ],
        [["GU-FG", "BLS-CHO-TB30", "B2312", "2026-06-01", "2027-12-12", "1200", "14250.5"],
         ["KSN-SINARJAYA", "BLS-CHO-TB30", "B2312", "2026-06-01", "2027-12-12", "48", "14250.5"]],
    ),
    Sheet(
        "Neraca Saldo Awal",
        "Saldo akun neraca per tanggal cut-over (akhir periode akuntansi)",
        "Finance & Accounting",
        "A-36, ACC-14",
        [
            c("Kode Akun", True, "Kode COA final", "Hanya akun neraca (1xxx–3xxx)", "1111"),
            c("Nama Akun", False, TEXT, "Untuk pengecekan", "Bank BCA Operasional", width=26),
            c("Debit", False, MONEY, "Isi salah satu Debit/Kredit", "250000000"),
            c("Kredit", False, MONEY, "", ""),
        ],
        [["1111", "Bank BCA Operasional", "250000000", ""], ["2101", "Utang Usaha", "", "48500000"]],
    ),
    Sheet(
        "Piutang Terbuka",
        "Invoice pelanggan/mitra yang belum lunas per tanggal cut-over",
        "Finance & Accounting",
        "U-10, ACC-14",
        [
            c("Nomor Invoice Lama", True, TEXT, "Unik", "INV-LAMA-0931"),
            c("Tanggal Invoice", True, DATE, "", "2026-12-15", width=14),
            c("Jatuh Tempo", True, DATE, "", "2027-01-14", width=14),
            c("Kode Pelanggan/Mitra", True, TEXT, "", "MT-SINARJAYA"),
            c("Channel", True, "Pilihan", "", "Konsinyasi",
              ["Modern Trade", "Konsinyasi", "Marketplace", "WA/IG", "B2B", "Karyawan"]),
            c("Nilai Invoice", True, MONEY, "", "4500000"),
            c("Sisa Belum Dibayar", True, MONEY, "<= Nilai Invoice", "4500000"),
        ],
        [["INV-LAMA-0931", "2026-12-15", "2027-01-14", "MT-SINARJAYA", "Konsinyasi", "4500000", "4500000"]],
    ),
    Sheet(
        "Utang Terbuka",
        "Tagihan pemasok/maklon yang belum dibayar per tanggal cut-over",
        "Finance & Accounting",
        "U-10, ACC-14",
        [
            c("Nomor Tagihan Pemasok", True, TEXT, "Unik per pemasok", "KJ/INV/2026/551"),
            c("Kode Pemasok", True, TEXT, "", "SUP-KEMASJAYA"),
            c("Tanggal Tagihan", True, DATE, "", "2026-12-20", width=14),
            c("Jatuh Tempo", True, DATE, "", "2027-01-19", width=14),
            c("Nilai Tagihan", True, MONEY, "", "48500000"),
            c("Sisa Belum Dibayar", True, MONEY, "", "48500000"),
        ],
        [["KJ/INV/2026/551", "SUP-KEMASJAYA", "2026-12-20", "2027-01-19", "48500000", "48500000"]],
    ),
    Sheet(
        "PO Terbuka",
        "Purchase Order yang belum diterima penuh per tanggal cut-over",
        "Logistik",
        "U-10",
        [
            c("Nomor PO Lama", True, TEXT, "", "PO-LAMA-0120"),
            c("Kode Pemasok", True, TEXT, "", "MKL-NUTRIPRO"),
            c("Tanggal PO", True, DATE, "", "2026-12-01", width=14),
            c("Jenis PO", True, "Pilihan", "", "Jasa Maklon", ["Packaging", "Jasa Maklon", "Reprocess", "Lainnya"]),
            c("SKU/Kode", True, TEXT, "", "BLS-CHO-TB30"),
            c("Qty Dipesan", True, QTY, "", "10000"),
            c("Qty Sudah Diterima", True, QTY, "<= Qty Dipesan", "4000"),
            c("Harga per Unit", True, MONEY, "", "9500"),
            c("Uang Muka Sudah Dibayar", False, MONEY, "", "20000000"),
        ],
        [["PO-LAMA-0120", "MKL-NUTRIPRO", "2026-12-01", "Jasa Maklon", "BLS-CHO-TB30", "10000", "4000", "9500", "20000000"]],
    ),
    Sheet(
        "Aset Tetap",
        "Register aset tetap termasuk mesin vending",
        "Finance & Accounting",
        "ACC-21, U-11",
        [
            c("Kode Aset", True, TEXT, "Unik", "AST-VM-001"),
            c("Nama Aset", True, TEXT, "", "Mesin Vending Lokasi A", width=26),
            c("Kelas Aset", True, "Pilihan", "", "Mesin Vending", ["Mesin Vending", "Peralatan Kantor & Komputer", "Kendaraan"]),
            c("Tanggal Perolehan", True, DATE, "", "2025-07-01", width=14),
            c("Harga Perolehan", True, MONEY, "> 0", "45000000"),
            c("Umur Manfaat (bulan)", True, "Angka bulat", "> 0", "48"),
            c("Nilai Sisa", False, MONEY, "", "0"),
            c("Akumulasi Penyusutan per Cut-over", True, MONEY, "<= Harga Perolehan", "16875000"),
            c("Lokasi / Mitra", False, TEXT, "Untuk mesin vending: kode mitra", "MT-SINARJAYA"),
        ],
        [["AST-VM-001", "Mesin Vending Lokasi A", "Mesin Vending", "2025-07-01", "45000000", "48", "0", "16875000", "MT-SINARJAYA"]],
    ),
]

WORKBOOKS = {
    "Template-Migrasi-Master-Data.xlsx": ("Template Migrasi — Master Data (F0-MIG-01)", MASTER),
    "Template-Migrasi-Harga.xlsx": ("Template Migrasi — Harga (F0-MIG-02)", PRICE),
    "Template-Migrasi-Saldo-Awal.xlsx": ("Template Migrasi — Saldo Awal per Tanggal Cut-over (F0-MIG-03)", OPENING),
}

GENERAL_RULES = [
    "Isi data mulai baris 2 di setiap sheet data; jangan mengubah nama sheet, nama kolom, atau urutan kolom.",
    "Kolom bertanda * (header oranye) wajib diisi; kolom berheader abu-abu boleh dikosongkan.",
    "Tanggal ditulis YYYY-MM-DD (mis. 2027-01-31). Angka tanpa titik ribuan dan tanpa 'Rp'; desimal memakai titik.",
    "Kode (SKU, kode mitra, kode pemasok, kode akun) ditulis persis sama antar sheet dan antar template.",
    "Kolom pilihan memakai daftar dropdown; nilai di luar daftar akan ditolak saat impor.",
    "Contoh baris ada di sheet Petunjuk; hapus semua data contoh sebelum menyerahkan file.",
    "Saat diimpor, sistem memvalidasi setiap baris dan mengembalikan file berisi baris yang error beserta keterangannya.",
]


def build_petunjuk(wb: Workbook, title: str, sheets: list[Sheet]) -> None:
    ws = wb.active
    ws.title = "Petunjuk"
    ws["A1"] = title
    ws["A1"].font = TITLE_FONT
    ws["A2"] = "Balans Nutrition — Unified Business Platform"
    row = 4
    ws.cell(row=row, column=1, value="Aturan umum").font = Font(bold=True)
    for rule in GENERAL_RULES:
        row += 1
        ws.cell(row=row, column=1, value=f"• {rule}")
    row += 2
    ws.cell(row=row, column=1, value="Ringkasan sheet").font = Font(bold=True)
    row += 1
    for i, h in enumerate(["Sheet", "Isi", "PIC Balans", "Requirement"], start=1):
        ws.cell(row=row, column=i, value=h).font = Font(bold=True)
    for s in sheets:
        row += 1
        for i, v in enumerate([s.title, s.purpose, s.pic, s.reference], start=1):
            ws.cell(row=row, column=i, value=v)
    for s in sheets:
        row += 2
        ws.cell(row=row, column=1, value=f"Kolom sheet '{s.title}'").font = Font(bold=True, size=12)
        row += 1
        for i, h in enumerate(["Kolom", "Wajib", "Format", "Aturan validasi", "Contoh"], start=1):
            ws.cell(row=row, column=i, value=h).font = Font(bold=True)
        for col in s.columns:
            row += 1
            values = [col.name, "Ya" if col.required else "Tidak", col.fmt,
                      col.rule + (f" Pilihan: {', '.join(col.choices)}." if col.choices else ""), col.example]
            for i, v in enumerate(values, start=1):
                ws.cell(row=row, column=i, value=v).alignment = Alignment(wrap_text=True, vertical="top")
        row += 1
        ws.cell(row=row, column=1, value="Contoh baris:").font = EXAMPLE_FONT
        for example in s.examples:
            row += 1
            for i, v in enumerate(example, start=1):
                ws.cell(row=row, column=i, value=v).font = EXAMPLE_FONT
    for col, width in zip("ABCDE", [30, 44, 30, 48, 24]):
        ws.column_dimensions[col].width = width


def build_data_sheet(wb: Workbook, s: Sheet) -> None:
    ws = wb.create_sheet(s.title)
    for i, col in enumerate(s.columns, start=1):
        cell = ws.cell(row=1, column=i, value=f"{col.name}{' *' if col.required else ''}")
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = HEADER_FILL if col.required else OPTIONAL_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        letter = get_column_letter(i)
        ws.column_dimensions[letter].width = col.width
        if col.choices:
            dv = DataValidation(type="list", formula1='"' + ",".join(col.choices) + '"', allow_blank=not col.required,
                                showErrorMessage=True, errorTitle="Nilai tidak valid",
                                error=f"Pilih salah satu: {', '.join(col.choices)}")
            dv.add(f"{letter}2:{letter}5000")
            ws.add_data_validation(dv)
        if "Tanggal" in col.fmt or col.fmt == DATE:
            for r in range(2, 5001):
                ws.cell(row=r, column=i).number_format = "@"
    ws.row_dimensions[1].height = 32
    ws.freeze_panes = "A2"


def main(out_dir: str) -> None:
    out = Path(out_dir)
    out.mkdir(parents=True, exist_ok=True)
    for filename, (title, sheets) in WORKBOOKS.items():
        wb = Workbook()
        build_petunjuk(wb, title, sheets)
        for s in sheets:
            build_data_sheet(wb, s)
        wb.save(out / filename)
        print(f"ditulis: {out / filename}")


if __name__ == "__main__":
    main(sys.argv[1] if len(sys.argv) > 1 else "docs/migrasi")
